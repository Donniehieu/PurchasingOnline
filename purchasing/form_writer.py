"""Điền kết quả vào form duyệt báo giá (dựa trên file mẫu, giữ nguyên tiêu đề và cấu trúc cột A-H).

Cột I 'Ngày cập nhật trên web' được thêm mới. Ngoài sheet chính, file còn có các sheet:
Thông tin yêu cầu, So sánh giá (biểu đồ), Đàm phán NCC, Thuyết phục sếp, Nguồn tham khảo, Hình ảnh.
"""

from copy import copy
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import insights

HIGHLIGHT_FILL = PatternFill("solid", fgColor="92D050")
HEADER_FILL = PatternFill("solid", fgColor="003399")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
FIRST_DATA_ROW = 3
DATE_COLUMN = 9
SOURCE_HEADERS = ["STT", "NCC", "Tiêu đề trang", "Link tham khảo", "Giá tìm được (đ)", "Ngày cập nhật web", "Đủ điều kiện", "Ghi chú"]


def _notes(item):
    lines = []
    if item.get("note"):
        lines.append(f"Ghi chú người mua: {item['note']}")
    if item.get("recommended"):
        lines.append(f"★ KHUYẾN NGHỊ ƯU TIÊN (điểm {item['score']})")
    if not item["eligible"]:
        lines.append(f"Loại: {item['reason']}")
    lines.append(f"Bảo hành: {item['warranty_months']} tháng" if item["warranty_months"] else "Bảo hành: chưa rõ, cần hỏi NCC")
    lines.append(f"Hóa đơn VAT: {_yes_no(item['vat'])}")
    if item["vat_included"]:
        lines.append("Giá trên web đã gồm VAT, đơn giá đã chia 1.1")
    if item.get("description"):
        lines.append(f"Mô tả: {item['description'][:160]}")
    phones = (item.get("contacts") or {}).get("phones") or []
    if phones:
        lines.append("Liên hệ: " + ", ".join(phones))
    lines.append("Trạng thái: cần xác minh trước khi đặt hàng")
    lines.append(f"Nguồn: {item['url']}")
    return "\n".join(lines)


def _yes_no(value):
    return {True: "Có", False: "Không", None: "Chưa rõ"}[value]


def _spec(item):
    parts = []
    if item["mah"]:
        parts.append(f"Pin {item['mah']:,} mAh".replace(",", "."))
    if item["volt"]:
        parts.append(f"{item['volt']}V")
    parts.append(f"Trang: {item['title']}")
    return ", ".join(parts)


def _as_date(iso):
    if not iso:
        return None
    return datetime.strptime(iso, "%Y-%m-%d").date()


def write_form(template_path, output_path, request_info, suppliers, sources, image_paths=()):
    """Tạo file báo giá mới từ template.

    request_info: product, product_code, quantity, purpose, keywords, engines, errors
    suppliers: danh sách đã chấm điểm (NCC đủ điều kiện đứng trước)
    sources: toàn bộ NCC đã đọc, kể cả bị loại
    """
    wb = load_workbook(template_path)
    ws = wb.worksheets[0]
    _clear_data_rows(ws)
    _write_header_extra(ws)

    quantity = request_info["quantity"]
    for idx, item in enumerate(suppliers):
        row = FIRST_DATA_ROW + idx
        price = item["price"]
        if price and item["vat_included"]:
            price = round(price / 1.1)
        ws.cell(row, 1, idx + 1)
        ws.cell(row, 4, item["name"])
        ws.cell(row, 6, _spec(item))
        ws.cell(row, 7, price)
        ws.cell(row, 8, _notes(item))
        ws.cell(row, DATE_COLUMN, _as_date(item.get("date")))
        ws.cell(row, 7).number_format = "#,##0"
        ws.cell(row, DATE_COLUMN).number_format = "dd/mm/yyyy"
        ws.row_dimensions[row].height = max(30, 15 * _line_count(ws.cell(row, 8).value))
        for col in range(1, DATE_COLUMN + 1):
            cell = ws.cell(row, col)
            cell.border = BORDER
            cell.alignment = WRAP
            if item.get("recommended"):
                cell.fill = HIGHLIGHT_FILL
                cell.font = Font(bold=True)

    last_row = FIRST_DATA_ROW + max(len(suppliers), 1) - 1
    if suppliers:
        ws.cell(FIRST_DATA_ROW, 2, request_info["product"])
        ws.cell(FIRST_DATA_ROW, 3, quantity)
        if last_row > FIRST_DATA_ROW:
            ws.merge_cells(start_row=FIRST_DATA_ROW, start_column=2, end_row=last_row, end_column=2)
            ws.merge_cells(start_row=FIRST_DATA_ROW, start_column=3, end_row=last_row, end_column=3)
        for col in (2, 3):
            cell = ws.cell(FIRST_DATA_ROW, col)
            cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
            cell.border = BORDER
    ws.cell(last_row + 2, 2, "Mục đích sử dụng:").font = Font(bold=True)
    ws.cell(last_row + 2, 3, request_info["purpose"] or "Chưa nhập")
    ws.column_dimensions[get_column_letter(DATE_COLUMN)].width = 20

    _write_request_sheet(wb, request_info, len(sources))
    _write_price_sheet(wb, suppliers)
    analysis = insights.build_insights(suppliers, quantity, request_info.get("ctx"))
    _write_negotiation_sheet(wb, analysis["negotiation"], request_info)
    _write_persuasion_sheet(wb, analysis["persuasion"])
    _write_sources_sheet(wb, sources)
    if request_info.get("edits"):
        _write_edits_sheet(wb, request_info["edits"])
    if image_paths:
        _write_images_sheet(wb, image_paths)

    wb.save(output_path)
    return output_path, analysis


def _clear_data_rows(ws):
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= FIRST_DATA_ROW:
            ws.unmerge_cells(str(merged))
    if ws.max_row >= FIRST_DATA_ROW:
        ws.delete_rows(FIRST_DATA_ROW, ws.max_row)


def _write_header_extra(ws):
    """Thêm cột I 'Ngày cập nhật trên web' với kiểu chữ như các tiêu đề còn lại."""
    header = ws.cell(2, DATE_COLUMN, "Ngày cập nhật trên web")
    header._style = copy(ws.cell(2, 8)._style)
    header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _line_count(text):
    return sum(max(1, len(line) // 60 + 1) for line in str(text or "").split("\n"))


def _write_request_sheet(wb, info, source_count):
    _replace_sheet(wb, "Thông tin yêu cầu")
    ws = wb["Thông tin yêu cầu"]
    rows = [
        ("Tên sản phẩm", info["product"]),
        ("Mã / model", info["product_code"] or "Không nhập"),
        ("Số lượng", info["quantity"]),
        ("Mục đích sử dụng", info["purpose"] or "Chưa nhập"),
        ("Thông số đã xác nhận", info.get("spec") or "Không yêu cầu"),
        ("Từ khóa tìm kiếm", info["keywords"]),
        ("Engine đã dùng", ", ".join(info["engines"]) or "Chỉ link NCC nhập tay"),
        ("Số trang NCC đã đọc", source_count),
        ("Ngày lập", date.today().strftime("%d/%m/%Y")),
    ]
    if info.get("errors"):
        rows.append(("Lỗi engine", "; ".join(f"{k}: {v}" for k, v in info["errors"].items())))
    for label, value in rows:
        ws.append([label, value])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = WRAP
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 80


def _write_price_sheet(wb, suppliers):
    _replace_sheet(wb, "So sánh giá")
    ws = wb["So sánh giá"]
    ws.append(["NCC", "Đơn giá (đ)", "Tỷ lệ so với rẻ nhất (%)", "Khuyến nghị"])
    _style_header(ws)
    rows = insights.price_comparison(suppliers)
    for row in rows:
        ws.append([row["name"], row["price"], row["ratio"], "Có" if row["recommended"] else ""])
    for r in range(2, len(rows) + 2):
        ws.cell(r, 2).number_format = "#,##0"
    if rows:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "So sánh đơn giá giữa các NCC"
        chart.y_axis.title = "Đơn giá (đ)"
        data = Reference(ws, min_col=2, min_row=1, max_row=len(rows) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 10, 20
        ws.add_chart(chart, "F2")
    _set_widths(ws, (34, 16, 24, 14))


def _write_negotiation_sheet(wb, negotiation, info):
    _replace_sheet(wb, "Đàm phán NCC")
    ws = wb["Đàm phán NCC"]
    if not negotiation:
        ws.append(["Chưa có NCC khuyến nghị để lập kịch bản đàm phán."])
    for item in negotiation:
        ws.append([f"NCC: {item['ncc']}"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append([f"Giá hiện tại: {_money(item['current_price'])}"])
        ws.append([f"Mục tiêu: {item['target_note']}"])
        ws.append([f"Phương án dự phòng: {item['backup']}"])
        for title, lines in (
            ("Điểm mạnh thương lượng (từ dữ liệu thu thập)", item["leverage"]),
            ("Câu hỏi cần làm rõ với NCC", item["questions"]),
            ("Chiến lược thuyết phục NCC", item["strategies"]),
            ("Các bước đàm phán theo quy trình mua hàng", item["steps"]),
        ):
            ws.append([])
            ws.append([title])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for line in lines:
                ws.append([f"- {line}"])
        if item["references"]:
            ws.append([])
            ws.append(["Căn cứ tham khảo trên web"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for ref in item["references"]:
                ws.append([f"- [{ref['topic']}] {ref['text']} ({ref['url']})"])
    _set_widths(ws, (130,))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_persuasion_sheet(wb, persuasion):
    _replace_sheet(wb, "Thuyết phục sếp")
    ws = wb["Thuyết phục sếp"]
    if not persuasion["title"]:
        ws.append(["Chưa có NCC khuyến nghị để lập luận thuyết phục."])
    else:
        ws.append([persuasion["title"]])
        ws.cell(1, 1).font = Font(bold=True, size=13)
        for section in persuasion["sections"]:
            ws.append([])
            ws.append([section["heading"]])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for point in section["points"]:
                ws.append([f"- {point}"])
        ws.append([])
        ws.append([persuasion["closing"]])
    _set_widths(ws, (130,))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_edits_sheet(wb, edits):
    _replace_sheet(wb, "Lịch sử điều chỉnh")
    ws = wb["Lịch sử điều chỉnh"]
    ws.append(["Thời gian", "NCC", "Trường", "Giá trị cũ", "Giá trị mới"])
    _style_header(ws)
    for edit in edits:
        ws.append([edit["time"].replace("T", " "), edit["supplier"], edit["field_label"], str(edit["old"]), str(edit["new"])])
    _set_widths(ws, (20, 30, 22, 30, 30))


def _money(value):
    return f"{round(value or 0):,} đ".replace(",", ".")


def _write_sources_sheet(wb, suppliers):
    _replace_sheet(wb, "Nguồn tham khảo")
    ws = wb["Nguồn tham khảo"]
    ws.append(SOURCE_HEADERS)
    _style_header(ws)
    for idx, item in enumerate(suppliers, start=1):
        ws.append([
            idx, item["name"], item["title"], item["url"], item["price"],
            _as_date(item.get("date")), "Có" if item.get("eligible") else "Không", _notes(item),
        ])
        ws.cell(idx + 1, 5).number_format = "#,##0"
        ws.cell(idx + 1, 6).number_format = "dd/mm/yyyy"
    _set_widths(ws, (6, 26, 40, 60, 18, 18, 14, 60))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_images_sheet(wb, image_paths):
    _replace_sheet(wb, "Hình ảnh")
    ws = wb["Hình ảnh"]
    anchor_row = 1
    for path in image_paths:
        img = XLImage(str(path))
        scale = 240 / img.width if img.width > 240 else 1
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, f"A{anchor_row}")
        anchor_row += max(img.height // 20, 1) + 2


def _replace_sheet(wb, title):
    if title in wb.sheetnames:
        del wb[title]
    wb.create_sheet(title)


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _set_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
