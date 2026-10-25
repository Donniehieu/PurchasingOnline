"""Tổng hợp lịch sử tìm kiếm theo tên sản phẩm: số lượt, từ khóa, trang web, NCC, giá trung bình, ..."""

import io
import statistics
import unicodedata
from collections import Counter
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COUNTED_STATUSES = ("done", "error", "running", "queued")


def _key(name):
    text = unicodedata.normalize("NFD", name or "")
    text = "".join(c for c in text if unicodedata.category(c) != "Mn").replace("đ", "d").lower()
    return " ".join(text.split())


def _domain(url):
    return urlparse(url).netloc.lower().removeprefix("www.")


def group_history(jobs):
    """jobs: iterable job.json. Trả về danh sách nhóm theo sản phẩm, nhóm tìm gần nhất đứng trước."""
    groups = {}
    for job in jobs:
        if job["status"] not in COUNTED_STATUSES or not job.get("query"):
            continue
        groups.setdefault(_key(job["query"]), []).append(job)

    result = []
    for key, items in groups.items():
        items.sort(key=lambda j: j.get("created", ""))
        keywords, engines, site_counter = [], set(), Counter()
        supplier_stats, recommended = {}, Counter()
        prices, job_means = [], []
        for job in items:
            for word in [job.get("keywords")] + [r.get("query") for r in job.get("research", {}).values()]:
                if word and word not in keywords:
                    keywords.append(word)
            engines.update(job.get("engines", []))
            for url in set(job.get("sites_searched", [])) | {s["url"] for s in job["suppliers"]}:
                site_counter[_domain(url)] += 1
            job_prices = [s["price"] for s in job["suppliers"] if s.get("price")]
            prices += job_prices
            if job_prices:
                job_means.append(statistics.mean(job_prices))
            for s in job["suppliers"]:
                stat = supplier_stats.setdefault(s["name"], {"name": s["name"], "times": 0, "min_price": None, "last_price": None, "url": s["url"]})
                stat["times"] += 1
                if s.get("price"):
                    stat["last_price"] = s["price"]
                    stat["min_price"] = s["price"] if stat["min_price"] is None else min(stat["min_price"], s["price"])
                if s.get("recommended"):
                    recommended[s["name"]] += 1
        change = None
        if len(job_means) >= 2 and job_means[0]:
            change = round((job_means[-1] - job_means[0]) / job_means[0] * 100, 1)
        result.append({
            "key": key,
            "name": items[-1]["query"],
            "searches": len(items),
            "keywords": keywords,
            "sites": [{"domain": d, "visits": n} for d, n in site_counter.most_common()],
            "suppliers": sorted(supplier_stats.values(), key=lambda s: -s["times"]),
            "supplier_count": len(supplier_stats),
            "avg_price": round(statistics.mean(prices)) if prices else None,
            "median_price": round(statistics.median(prices)) if prices else None,
            "min_price": min(prices) if prices else None,
            "max_price": max(prices) if prices else None,
            "price_change_pct": change,
            "first_search": items[0].get("created", ""),
            "last_search": items[-1].get("created", ""),
            "engines": sorted(engines),
            "avg_suppliers_per_search": round(statistics.mean(len(j["suppliers"]) for j in items), 1),
            "total_quantity": sum(j.get("quantity") or 0 for j in items),
            "recommended": [{"name": n, "times": c} for n, c in recommended.most_common()],
            "edit_count": sum(len(j.get("edits", [])) for j in items),
            "jobs": [{"id": j["id"], "status": j["status"], "created": j.get("created", ""), "quantity": j.get("quantity")} for j in reversed(items)],
        })
    return sorted(result, key=lambda g: g["last_search"], reverse=True)


def history_to_xlsx(groups):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp theo sản phẩm"
    headers = ["Sản phẩm", "Số lượt tìm", "Từ khóa đã dùng", "Số trang web", "Các trang web", "Số NCC", "Các NCC",
               "Giá trung bình (đ)", "Giá trung vị (đ)", "Giá thấp nhất (đ)", "Giá cao nhất (đ)", "Biến động giá (%)",
               "Lần đầu tìm", "Lần gần nhất", "Engine đã dùng", "NCC được khuyến nghị", "Số lần sửa bảng"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003399")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for g in groups:
        ws.append([
            g["name"], g["searches"], "\n".join(g["keywords"]), len(g["sites"]), "\n".join(s["domain"] for s in g["sites"]),
            g["supplier_count"], "\n".join(s["name"] for s in g["suppliers"]), g["avg_price"], g["median_price"],
            g["min_price"], g["max_price"], g["price_change_pct"], g["first_search"].replace("T", " "),
            g["last_search"].replace("T", " "), ", ".join(g["engines"]),
            ", ".join(f"{r['name']} ({r['times']})" for r in g["recommended"]), g["edit_count"],
        ])
    for col, width in zip("ABCDEFGHIJKLMNOPQ", (30, 10, 45, 10, 40, 8, 40, 16, 16, 16, 16, 12, 20, 20, 18, 36, 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (7, 8, 9, 10):
            row[idx].number_format = "#,##0"

    detail = wb.create_sheet("Chi tiết NCC")
    detail.append(["Sản phẩm", "NCC", "Số lần xuất hiện", "Giá thấp nhất (đ)", "Giá gần nhất (đ)", "Link"])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for g in groups:
        for s in g["suppliers"]:
            detail.append([g["name"], s["name"], s["times"], s["min_price"], s["last_price"], s["url"]])
    for col, width in zip("ABCDEF", (30, 30, 14, 18, 18, 80)):
        detail.column_dimensions[col].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def edits_to_xlsx(edits):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lịch sử điều chỉnh"
    ws.append(["Thời gian", "Sản phẩm", "NCC", "Trường", "Giá trị cũ", "Giá trị mới", "Mã yêu cầu"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for e in edits:
        ws.append([e["time"].replace("T", " "), e["product"], e["supplier"], e["field_label"], str(e["old"]), str(e["new"]), e["job_id"]])
    for col, width in zip("ABCDEFG", (20, 30, 28, 22, 28, 28, 14)):
        ws.column_dimensions[col].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
