"""Nhập / xuất danh sách link NCC bằng file Excel (.xlsx) hoặc CSV, hoặc .txt mỗi dòng một link."""

import csv
import io

from openpyxl import Workbook, load_workbook


def parse_link_file(filename, data):
    """Lấy mọi ô có dạng http(s)://... từ file, giữ thứ tự và bỏ trùng."""
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        values = _values_from_xlsx(data)
    else:
        values = [cell for row in csv.reader(io.StringIO(data.decode("utf-8-sig"))) for cell in row]
    urls = []
    for value in values:
        text = str(value or "").strip()
        for part in text.split():
            if part.startswith(("http://", "https://")) and part not in urls:
                urls.append(part)
    return urls


def _values_from_xlsx(data):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                yield value


def links_to_csv(rows):
    """rows: danh sách (tên NCC, link) -> chuỗi CSV để xuất file."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["NCC", "Link"])
    writer.writerows(rows)
    return buffer.getvalue()


def links_to_xlsx(rows):
    """rows: danh sách (tên NCC, link) -> nội dung file .xlsx dạng bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Link NCC"
    ws.append(["NCC", "Link"])
    for row in rows:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 90
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
